"""
File Converter API Server
Standalone API server for Qwen integration (no GUI required)
Converts various formats (md, json, html, csv, py) to docx, xlsx, pptx
"""

import os
import sys
import json
import tempfile
import base64
import re
import csv
import io
from typing import Optional

# Web server imports
from flask import Flask, request, jsonify
from flask_cors import CORS

# Document processing imports
try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    Document = None

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError:
    Presentation = None


class FileConverter:
    """Core conversion logic for various file formats"""
    
    @staticmethod
    def md_to_docx(content: str, output_path: str) -> str:
        """Convert Markdown to DOCX"""
        if Document is None:
            raise ImportError("python-docx not installed")
        
        doc = Document()
        lines = content.split('\n')
        current_paragraph = []
        
        for line in lines:
            stripped = line.strip()
            
            if stripped.startswith('# '):
                if current_paragraph:
                    doc.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith('## '):
                if current_paragraph:
                    doc.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith('### '):
                if current_paragraph:
                    doc.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith('```'):
                if current_paragraph:
                    doc.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
            elif stripped.startswith('- ') or stripped.startswith('* '):
                if current_paragraph:
                    doc.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                doc.add_paragraph(stripped[2:], style='List Bullet')
            elif stripped:
                current_paragraph.append(stripped)
            else:
                if current_paragraph:
                    doc.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                doc.add_paragraph('')
        
        if current_paragraph:
            doc.add_paragraph(' '.join(current_paragraph))
        
        doc.save(output_path)
        return output_path
    
    @staticmethod
    def json_to_docx(content: str, output_path: str) -> str:
        """Convert JSON to DOCX"""
        if Document is None:
            raise ImportError("python-docx not installed")
        
        data = json.loads(content)
        doc = Document()
        doc.add_heading('JSON Data', level=1)
        
        def add_json_data(data, indent=0):
            if isinstance(data, dict):
                for key, value in data.items():
                    p = doc.add_paragraph()
                    runner = p.add_run(f"{'  ' * indent}{key}: ")
                    runner.bold = True
                    add_json_data(value, indent + 1)
            elif isinstance(data, list):
                for i, item in enumerate(data):
                    p = doc.add_paragraph()
                    runner = p.add_run(f"{'  ' * indent}[{i}]: ")
                    add_json_data(item, indent + 1)
            else:
                doc.add_paragraph(f"{'  ' * indent}{data}")
        
        add_json_data(data)
        doc.save(output_path)
        return output_path
    
    @staticmethod
    def html_to_docx(content: str, output_path: str) -> str:
        """Convert HTML to DOCX"""
        if Document is None:
            raise ImportError("python-docx not installed")
        
        doc = Document()
        content = re.sub(r'<script[^>]*>.*?</script>', '', content, flags=re.DOTALL | re.IGNORECASE)
        content = re.sub(r'<style[^>]*>.*?</style>', '', content, flags=re.DOTALL | re.IGNORECASE)
        content = html.unescape(content) if hasattr(html, 'unescape') else content
        
        blocks = re.split(r'</(?:p|div|h[1-6]|br|li)?>', content, flags=re.IGNORECASE)
        for block in blocks:
            text = re.sub(r'<[^>]+>', ' ', block).strip()
            if text:
                doc.add_paragraph(text)
        
        doc.save(output_path)
        return output_path
    
    @staticmethod
    def py_to_docx(content: str, output_path: str) -> str:
        """Convert Python code to DOCX"""
        if Document is None:
            raise ImportError("python-docx not installed")
        
        doc = Document()
        doc.add_heading('Python Code', level=1)
        p = doc.add_paragraph()
        runner = p.add_run(content)
        runner.font.name = 'Courier New'
        runner.font.size = Pt(10)
        doc.save(output_path)
        return output_path
    
    @staticmethod
    def md_to_pptx(content: str, output_path: str) -> str:
        """Convert Markdown to PPTX"""
        if Presentation is None:
            raise ImportError("python-pptx not installed")
        
        prs = Presentation()
        lines = content.split('\n')
        current_slide_content = []
        slide_title = "Slide"
        
        def create_slide(title, content_lines):
            if len(prs.slides) == 0:
                layout = prs.slide_layouts[0]
                slide = prs.slides.add_slide(layout)
                slide.shapes.title.text = title
                if content_lines:
                    subtitle = slide.placeholders[1]
                    subtitle.text = '\n'.join(content_lines[:5])
            else:
                layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(layout)
                slide.shapes.title.text = title
                body = slide.placeholders[1]
                tf = body.text_frame
                tf.clear()
                for i, line in enumerate(content_lines):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = line
        
        for line in lines:
            stripped = line.strip()
            if stripped.startswith('# '):
                if current_slide_content:
                    create_slide(slide_title, current_slide_content)
                    current_slide_content = []
                slide_title = stripped[2:]
            elif stripped.startswith('## '):
                if current_slide_content:
                    create_slide(slide_title, current_slide_content)
                    current_slide_content = []
                slide_title = stripped[3:]
            elif stripped.startswith('- ') or stripped.startswith('* '):
                current_slide_content.append(stripped[2:])
            elif stripped:
                current_slide_content.append(stripped)
        
        if current_slide_content or slide_title != "Slide":
            create_slide(slide_title, current_slide_content)
        
        if len(prs.slides) == 0:
            layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = "Presentation"
        
        prs.save(output_path)
        return output_path
    
    @staticmethod
    def json_to_pptx(content: str, output_path: str) -> str:
        """Convert JSON to PPTX"""
        if Presentation is None:
            raise ImportError("python-pptx not installed")
        
        data = json.loads(content)
        prs = Presentation()
        
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "JSON Data Presentation"
        
        def add_data_to_slides(data, title="Data"):
            if isinstance(data, dict):
                layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(layout)
                slide.shapes.title.text = title
                body = slide.placeholders[1]
                tf = body.text_frame
                tf.clear()
                for i, (key, value) in enumerate(data.items()):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = f"{key}: {value}"
                    if isinstance(value, (dict, list)):
                        add_data_to_slides(value, key)
            elif isinstance(data, list):
                layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(layout)
                slide.shapes.title.text = title
                body = slide.placeholders[1]
                tf = body.text_frame
                tf.clear()
                for i, item in enumerate(data):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = f"Item {i}: {item}"
        
        if isinstance(data, dict):
            add_data_to_slides(data)
        elif isinstance(data, list):
            add_data_to_slides(data, "List Items")
        
        prs.save(output_path)
        return output_path
    
    @staticmethod
    def html_to_pptx(content: str, output_path: str) -> str:
        """Convert HTML to PPTX"""
        if Presentation is None:
            raise ImportError("python-pptx not installed")
        
        prs = Presentation()
        content = re.sub(r'<script[^>]*>.*?</script>', '', content, flags=re.DOTALL | re.IGNORECASE)
        content = re.sub(r'<style[^>]*>.*?</style>', '', content, flags=re.DOTALL | re.IGNORECASE)
        content = html.unescape(content) if hasattr(html, 'unescape') else content
        
        headings = re.findall(r'<h[1-6][^>]*>(.*?)</h[1-6]>', content, re.IGNORECASE)
        paragraphs = re.findall(r'<p[^>]*>(.*?)</p>', content, re.IGNORECASE)
        
        if headings:
            for i, heading in enumerate(headings):
                heading = re.sub(r'<[^>]+>', '', heading).strip()
                slide_content = []
                if i < len(paragraphs):
                    para = re.sub(r'<[^>]+>', '', paragraphs[i]).strip()
                    slide_content.append(para)
                
                if i == 0:
                    layout = prs.slide_layouts[0]
                    slide = prs.slides.add_slide(layout)
                    slide.shapes.title.text = heading
                else:
                    layout = prs.slide_layouts[1]
                    slide = prs.slides.add_slide(layout)
                    slide.shapes.title.text = heading
                    if slide_content:
                        body = slide.placeholders[1]
                        body.text_frame.paragraphs[0].text = slide_content[0]
        else:
            layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = "HTML Content"
            if paragraphs:
                body = slide.placeholders[1]
                body.text = re.sub(r'<[^>]+>', '', paragraphs[0]).strip()
        
        prs.save(output_path)
        return output_path
    
    @staticmethod
    def py_to_pptx(content: str, output_path: str) -> str:
        """Convert Python code to PPTX"""
        if Presentation is None:
            raise ImportError("python-pptx not installed")
        
        prs = Presentation()
        
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "Python Code"
        slide.placeholders[1].text = "Code Presentation"
        
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "Source Code"
        
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()
        
        lines = content.split('\n')[:20]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
        
        prs.save(output_path)
        return output_path
    
    @staticmethod
    def csv_to_xlsx(content: str, output_path: str) -> str:
        """Convert CSV to XLSX"""
        if Workbook is None:
            raise ImportError("openpyxl not installed")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        
        reader = csv.reader(io.StringIO(content))
        for row_idx, row in enumerate(reader, 1):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column].width = min(max_length + 2, 50)
        
        wb.save(output_path)
        return output_path
    
    @staticmethod
    def json_to_xlsx(content: str, output_path: str) -> str:
        """Convert JSON to XLSX"""
        if Workbook is None:
            raise ImportError("openpyxl not installed")
        
        data = json.loads(content)
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        
        def flatten_json(data, prefix='', result=None):
            if result is None:
                result = {}
            if isinstance(data, dict):
                for key, value in data.items():
                    new_key = f"{prefix}.{key}" if prefix else key
                    flatten_json(value, new_key, result)
            elif isinstance(data, list):
                for i, item in enumerate(data):
                    flatten_json(item, f"{prefix}[{i}]", result)
            else:
                result[prefix] = data
            return result
        
        if isinstance(data, list):
            flattened_data = [flatten_json(item) for item in data]
            all_keys = set()
            for item in flattened_data:
                all_keys.update(item.keys())
            
            headers = sorted(list(all_keys))
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=1, column=col_idx, value=header)
            
            for row_idx, item in enumerate(flattened_data, 2):
                for col_idx, header in enumerate(headers, 1):
                    ws.cell(row=row_idx, column=col_idx, value=item.get(header, ''))
        else:
            flattened = flatten_json(data)
            for row_idx, (key, value) in enumerate(flattened.items(), 1):
                ws.cell(row=row_idx, column=1, value=key)
                ws.cell(row=row_idx, column=2, value=value)
        
        wb.save(output_path)
        return output_path
    
    @staticmethod
    def html_to_xlsx(content: str, output_path: str) -> str:
        """Convert HTML to XLSX"""
        if Workbook is None:
            raise ImportError("openpyxl not installed")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "HTML Data"
        
        tables = re.findall(r'<table[^>]*>(.*?)</table>', content, re.DOTALL | re.IGNORECASE)
        
        if tables:
            table = tables[0]
            rows = re.findall(r'<tr[^>]*>(.*?)</tr>', table, re.DOTALL | re.IGNORECASE)
            for row_idx, row in enumerate(rows, 1):
                cells = re.findall(r'<t[dh][^>]*>(.*?)</t[dh]>', row, re.DOTALL | re.IGNORECASE)
                for col_idx, cell in enumerate(cells, 1):
                    cell_text = re.sub(r'<[^>]+>', '', cell).strip()
                    cell_text = html.unescape(cell_text) if hasattr(html, 'unescape') else cell_text
                    ws.cell(row=row_idx, column=col_idx, value=cell_text)
        else:
            content = re.sub(r'<script[^>]*>.*?</script>', '', content, flags=re.DOTALL | re.IGNORECASE)
            content = re.sub(r'<style[^>]*>.*?</style>', '', content, flags=re.DOTALL | re.IGNORECASE)
            content = html.unescape(content) if hasattr(html, 'unescape') else content
            text = re.sub(r'<[^>]+>', '\n', content)
            lines = text.split('\n')
            for row_idx, line in enumerate(lines[:100], 1):
                if line.strip():
                    ws.cell(row=row_idx, column=1, value=line.strip())
        
        wb.save(output_path)
        return output_path
    
    @staticmethod
    def py_to_xlsx(content: str, output_path: str) -> str:
        """Convert Python code to XLSX"""
        if Workbook is None:
            raise ImportError("openpyxl not installed")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Python Code"
        
        ws.cell(row=1, column=1, value="Line Number")
        ws.cell(row=1, column=2, value="Code")
        
        lines = content.split('\n')
        for row_idx, line in enumerate(lines, 2):
            ws.cell(row=row_idx, column=1, value=row_idx - 1)
            ws.cell(row=row_idx, column=2, value=line)
        
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 80
        
        wb.save(output_path)
        return output_path


def is_conversion_supported(source: str, target: str) -> bool:
    """Check if conversion is supported"""
    supported = {
        'docx': ['md', 'json', 'html', 'py'],
        'pptx': ['md', 'json', 'html', 'py'],
        'xlsx': ['csv', 'json', 'html', 'py']
    }
    return source in supported.get(target, [])


def create_app():
    """Create Flask application"""
    app = Flask(__name__)
    CORS(app)
    converter = FileConverter()
    
    @app.route('/api/convert', methods=['POST'])
    def api_convert():
        try:
            data = request.get_json()
            
            if not data:
                return jsonify({"error": "No JSON data provided"}), 400
            
            content = data.get('content', '')
            source_format = data.get('source_format', 'md')
            target_format = data.get('target_format', 'docx')
            filename = data.get('filename', 'converted')
            
            if not content:
                return jsonify({"error": "No content provided"}), 400
            
            if not is_conversion_supported(source_format, target_format):
                return jsonify({
                    "error": f"Unsupported conversion: {source_format} to {target_format}",
                    "supported": {
                        "docx": ["md", "json", "html", "py"],
                        "pptx": ["md", "json", "html", "py"],
                        "xlsx": ["csv", "json", "html", "py"]
                    }
                }), 400
            
            temp_dir = tempfile.gettempdir()
            output_filename = f"{filename}.{target_format}"
            output_path = os.path.join(temp_dir, output_filename)
            
            # Perform conversion
            if target_format == "docx":
                if source_format == "md":
                    converter.md_to_docx(content, output_path)
                elif source_format == "json":
                    converter.json_to_docx(content, output_path)
                elif source_format == "html":
                    converter.html_to_docx(content, output_path)
                elif source_format == "py":
                    converter.py_to_docx(content, output_path)
            
            elif target_format == "pptx":
                if source_format == "md":
                    converter.md_to_pptx(content, output_path)
                elif source_format == "json":
                    converter.json_to_pptx(content, output_path)
                elif source_format == "html":
                    converter.html_to_pptx(content, output_path)
                elif source_format == "py":
                    converter.py_to_pptx(content, output_path)
            
            elif target_format == "xlsx":
                if source_format == "csv":
                    converter.csv_to_xlsx(content, output_path)
                elif source_format == "json":
                    converter.json_to_xlsx(content, output_path)
                elif source_format == "html":
                    converter.html_to_xlsx(content, output_path)
                elif source_format == "py":
                    converter.py_to_xlsx(content, output_path)
            
            # Read file and encode as base64
            with open(output_path, 'rb') as f:
                file_content = base64.b64encode(f.read()).decode('utf-8')
            
            # Clean up
            os.remove(output_path)
            
            return jsonify({
                "success": True,
                "filename": output_filename,
                "content_base64": file_content,
                "message": "Conversion successful"
            })
            
        except json.JSONDecodeError as e:
            return jsonify({"error": f"Invalid JSON: {e}"}), 400
        except Exception as e:
            return jsonify({"error": str(e)}), 500
    
    @app.route('/api/health', methods=['GET'])
    def health_check():
        return jsonify({
            "status": "healthy",
            "service": "file-converter-api",
            "version": "1.0.0"
        })
    
    @app.route('/', methods=['GET'])
    def index():
        return jsonify({
            "service": "File Converter API",
            "description": "Convert MD/JSON/HTML/CSV/PY files to DOCX/XLSX/PPTX",
            "endpoints": {
                "POST /api/convert": "Convert content to target format",
                "GET /api/health": "Health check"
            },
            "usage": {
                "content": "string (required) - The content to convert",
                "source_format": "string (optional) - Source format: md, json, html, csv, py",
                "target_format": "string (optional) - Target format: docx, xlsx, pptx",
                "filename": "string (optional) - Output filename (without extension)"
            },
            "example_request": {
                "content": "# Hello\nThis is markdown",
                "source_format": "md",
                "target_format": "docx",
                "filename": "my_document"
            }
        })
    
    return app


def main():
    """Main entry point for API server"""
    import argparse
    
    parser = argparse.ArgumentParser(description='File Converter API Server')
    parser.add_argument('--host', default='0.0.0.0', help='Host to bind to')
    parser.add_argument('--port', type=int, default=5000, help='Port to listen on')
    parser.add_argument('--debug', action='store_true', help='Enable debug mode')
    
    args = parser.parse_args()
    
    app = create_app()
    print(f"Starting File Converter API Server on http://{args.host}:{args.port}")
    print("API Endpoints:")
    print("  GET  /           - API documentation")
    print("  POST /api/convert - Convert files")
    print("  GET  /api/health  - Health check")
    print("\nQwen can use this API by sending POST requests to /api/convert")
    
    app.run(host=args.host, port=args.port, debug=args.debug)


if __name__ == "__main__":
    main()
