#!/usr/bin/env python3
"""
HTML Converter Pro - Python Backend
Конвертация HTML в DOCX, XLSX, PPTX с использованием Python библиотек
python-docx, python-pptx, openpyxl
"""

import sys
import json
import base64
from io import BytesIO
from pathlib import Path

# Импорты библиотек
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.enum.text import PP_ALIGN
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def parse_html_to_blocks(html_content):
    """Парсинг HTML в структуру блоков (упрощенная версия)"""
    from html.parser import HTMLParser
    
    blocks = []
    current_text = ""
    current_tag = None
    list_items = []
    in_list = False
    list_ordered = False
    
    class HTMLBlockParser(HTMLParser):
        def handle_starttag(self, tag, attrs):
            nonlocal current_tag, current_text, in_list, list_ordered
            
            if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
                if current_text.strip():
                    blocks.append({'type': 'paragraph', 'text': current_text.strip()})
                current_text = ""
                current_tag = tag
            elif tag == 'p':
                if current_text.strip() and not current_tag:
                    pass
                current_text = ""
                current_tag = 'p'
            elif tag in ['ul', 'ol']:
                in_list = True
                list_ordered = (tag == 'ol')
                list_items = []
            elif tag == 'li':
                current_text = ""
            elif tag == 'br':
                current_text += "\n"
            elif tag == 'hr':
                blocks.append({'type': 'hr'})
            elif tag == 'blockquote':
                if current_text.strip():
                    blocks.append({'type': 'paragraph', 'text': current_text.strip()})
                current_text = ""
                current_tag = 'blockquote'
            elif tag == 'pre' or tag == 'code':
                if current_text.strip():
                    blocks.append({'type': 'paragraph', 'text': current_text.strip()})
                current_text = ""
                current_tag = 'code'
            elif tag == 'table':
                pass
        
        def handle_endtag(self, tag):
            nonlocal current_tag, current_text, in_list
            
            if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
                level = int(tag[1])
                blocks.append({
                    'type': 'heading',
                    'level': level,
                    'text': current_text.strip()
                })
                current_text = ""
                current_tag = None
            elif tag == 'p':
                if current_text.strip():
                    blocks.append({'type': 'paragraph', 'text': current_text.strip()})
                current_text = ""
                current_tag = None
            elif tag in ['ul', 'ol']:
                if list_items:
                    blocks.append({
                        'type': 'list',
                        'ordered': list_ordered,
                        'items': list_items
                    })
                list_items = []
                in_list = False
            elif tag == 'li':
                if current_text.strip():
                    list_items.append(current_text.strip())
                current_text = ""
            elif tag == 'blockquote':
                if current_text.strip():
                    blocks.append({
                        'type': 'quote',
                        'text': current_text.strip()
                    })
                current_text = ""
                current_tag = None
            elif tag in ['pre', 'code']:
                if current_text.strip():
                    blocks.append({
                        'type': 'code',
                        'text': current_text.strip()
                    })
                current_text = ""
                current_tag = None
        
        def handle_data(self, data):
            nonlocal current_text
            current_text += data
    
    parser = HTMLBlockParser()
    parser.feed(html_content)
    
    # Добавляем оставшийся текст
    if current_text.strip():
        if current_tag == 'p':
            blocks.append({'type': 'paragraph', 'text': current_text.strip()})
        elif not current_tag:
            blocks.append({'type': 'paragraph', 'text': current_text.strip()})
    
    return blocks


def generate_docx_python(html_content, output_path):
    """Генерация DOCX файла с использованием python-docx"""
    doc = DocxDocument()
    blocks = parse_html_to_blocks(html_content)
    
    for block in blocks:
        if block['type'] == 'heading':
            level = block['level']
            style_name = f'Heading {min(level, 9)}'
            try:
                p = doc.add_heading(block['text'], level=level)
            except:
                p = doc.add_paragraph(block['text'])
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(24 - level * 2)
        
        elif block['type'] == 'paragraph':
            doc.add_paragraph(block['text'])
        
        elif block['type'] == 'list':
            for i, item in enumerate(block['items']):
                prefix = f"{i + 1}. " if block['ordered'] else "• "
                doc.add_paragraph(f"{prefix}{item}")
        
        elif block['type'] == 'quote':
            p = doc.add_paragraph(block['text'])
            p.italic = True
            p.paragraph_format.left_indent = Cm(1)
        
        elif block['type'] == 'code':
            p = doc.add_paragraph(block['text'])
            p.runs[0].font.name = 'Courier New'
            p.runs[0].font.size = Pt(10)
            p.paragraph_format.shading.background_color = 'F5F5F5'
        
        elif block['type'] == 'hr':
            doc.add_paragraph('─' * 50)
    
    doc.save(output_path)
    return output_path


def generate_xlsx_python(html_content, output_path):
    """Генерация XLSX файла с использованием openpyxl"""
    wb = Workbook()
    ws = wb.active
    ws.title = "HTML Content"
    
    blocks = parse_html_to_blocks(html_content)
    row = 1
    
    for block in blocks:
        text = ""
        if block['type'] == 'heading':
            level = block['level']
            prefix = '#' * level + ' '
            text = prefix + block['text']
        elif block['type'] == 'paragraph':
            text = block['text']
        elif block['type'] == 'list':
            for i, item in enumerate(block['items']):
                prefix = f"{i + 1}. " if block['ordered'] else "- "
                ws.cell(row=row, column=1, value=prefix + item)
                row += 1
            continue
        elif block['type'] == 'quote':
            text = '> ' + block['text']
        elif block['type'] == 'code':
            text = block['text']
        
        if text:
            ws.cell(row=row, column=1, value=text)
            row += 1
    
    # Устанавливаем ширину колонки
    ws.column_dimensions['A'].width = 60
    
    wb.save(output_path)
    return output_path


def generate_pptx_python(html_content, output_path):
    """Генерация PPTX файла с использованием python-pptx"""
    prs = Presentation()
    blocks = parse_html_to_blocks(html_content)
    
    # Группируем блоки по слайдам
    slides_content = []
    current_slide = []
    
    for block in blocks:
        if block['type'] == 'heading' and block['level'] <= 2:
            if current_slide:
                slides_content.append(current_slide)
                current_slide = []
        current_slide.append(block)
        
        if len(current_slide) >= 6:
            slides_content.append(current_slide)
            current_slide = []
    
    if current_slide:
        slides_content.append(current_slide)
    
    if not slides_content:
        slides_content = [[{'type': 'paragraph', 'text': 'Пустая презентация'}]]
    
    for idx, slide_blocks in enumerate(slides_content):
        # Определяем тип слайда
        title_block = next((b for b in slide_blocks if b['type'] == 'heading' and b['level'] <= 2), None)
        
        if title_block:
            layout = prs.slide_layouts[1]  # Title and Content
        else:
            layout = prs.slide_layouts[1]
        
        slide = prs.slides.add_slide(layout)
        
        # Добавляем заголовок
        if title_block:
            title = slide.shapes.title
            title.text = title_block['text']
        
        # Добавляем контент
        content_y = 1.5 if title_block else 0.5
        for block in slide_blocks:
            if block == title_block:
                continue
            
            if block['type'] in ['paragraph', 'heading']:
                text = block.get('text', '')
                if block['type'] == 'heading':
                    text = f"{'#' * block['level']} {text}"
                
                # Находим placeholder для контента
                for shape in slide.shapes:
                    if hasattr(shape, "text_frame"):
                        tf = shape.text_frame
                        if tf.text == '' or tf.text == 'Click to add text':
                            p = tf.paragraphs[0] if tf.paragraphs else tf.add_paragraph()
                            p.text = text
                            p.font.size = PptxPt(16)
                            break
                break  # Добавляем только первый параграф для простоты
    
    prs.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 4:
        print("Usage: python html_converter.py <format> <html_file> <output_file>")
        print("Formats: docx, xlsx, pptx")
        sys.exit(1)
    
    format_type = sys.argv[1]
    html_file = sys.argv[2]
    output_file = sys.argv[3]
    
    # Читаем HTML файл
    with open(html_file, 'r', encoding='utf-8') as f:
        html_content = f.read()
    
    # Конвертируем
    if format_type == 'docx':
        result = generate_docx_python(html_content, output_file)
    elif format_type == 'xlsx':
        result = generate_xlsx_python(html_content, output_file)
    elif format_type == 'pptx':
        result = generate_pptx_python(html_content, output_file)
    else:
        print(f"Unknown format: {format_type}")
        sys.exit(1)
    
    print(f"Successfully created: {result}")


if __name__ == '__main__':
    main()
